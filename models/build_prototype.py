"""Build the spreadsheet prototype of the rain clause (report/prototipo.xlsx) and its HTML snapshot.

Sheet 1 ("Libro del vendedor") is the core calculation of the mechanism: one cart vendor's ledger
over 30 selling days with good days, rain-trigger days and light-rain days, run on two paths at
once. On a rain day the vendor first keeps up to a household floor (input B8) and the rest of that
day's sales goes to his debts, on both paths. (a) With the rain clause, the unpaid fiado is repaid in
thirds with no surcharge; (b) without it, the gap is covered with a gota a gota loan at a 20% flat
charge in daily installments, and a later rain day forces a new loan that also pays the old one's
installment.

Sheet 2 ("Numeros del negocio") holds the unit economics of the SAS and the payer's view with the
same assumptions as models/business_model.py.

Every output cell is a live Excel formula over the yellow input cells. The script then evaluates
the workbook with the `formulas` package, checks the results against an independent Python mirror
and against business_model.py, and writes report/prototipo-snapshot.html from the evaluated values.

User-facing labels are in Spanish because the report is in Spanish. Run with the scratchpad venv:
<venv>/bin/python models/build_prototype.py
"""

import html
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import business_model as bm  # noqa: E402

REPORT_DIR = os.path.join(os.path.dirname(HERE), "report")
XLSX_PATH = os.path.join(REPORT_DIR, "prototipo.xlsx")
HTML_PATH = os.path.join(REPORT_DIR, "prototipo-snapshot.html")

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
OUTPUT_FILL = PatternFill("solid", fgColor="E2EFDA")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
RAIN_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)
WHITE_BOLD = Font(bold=True, color="FFFFFF")
TITLE = Font(bold=True, size=14)
THIN = Side(style="thin", color="B8B8B8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PESOS = '"$"#,##0'
PCT = "0%"
PCT1 = "0.0%"

# Illustrative rain in mm between 7 a.m. and 7 p.m. for 30 selling days. Invented to show the
# calculation (3 trigger days, one of them a 2-day streak, plus 2 light-rain days). The real
# frequency is 26 trigger days in 294 selling days (own count on IDEAM station 0026055120).
EXAMPLE_RAIN_MM = {5: 12.4, 9: 1.2, 13: 7.0, 14: 9.2, 21: 3.4}
LEDGER_DAYS = 30
SNAPSHOT_DAYS = (1, 4, 5, 6, 7, 8, 9, 12, 13, 14, 15, 16, 17, 21, 30)  # compact view for the annex

# ------------------------------------------------------------------------------------------
# Sheet 1: vendor ledger
# ------------------------------------------------------------------------------------------
LEDGER_INPUTS = [
    # (row, label, value, number format, source)
    (4, "Fiado que saca cada mañana", 62_000, PESOS,
     "Cota superior: ventas diarias del ambulante menos su ingreso (DANE, EMICRON ambulantes 2025). No medido para carritos [S]"),
    (5, "Ingreso neto de un día normal", bm.VENDOR_INCOME_PER_SELLING_DAY, PESOS,
     "DANE, EMICRON ambulantes 2025, C. 24: unos $1,0 millón al mes en 26 días de venta"),
    (6, "Ventas en día de lluvia (% de un día normal)", 0.30, PCT, "[S] sin dato; primera pregunta del piloto"),
    (7, "Ventas en día de lluvia leve, bajo el umbral (%)", 0.85, PCT, "[S] sin dato"),
    (8, "Lo que se queda para la casa el día de lluvia (tope)", 15_000, PESOS,
     "[S] piso para la casa: de lo vendido ese día guarda hasta este monto y el resto abona a su deuda, "
     "igual en los dos caminos"),
    (9, "Umbral de lluvia (mm entre 7 a. m. y 7 p. m.)", 5, "0.0",
     "Regla de la cláusula; convención del frente 12b, a calibrar [S]"),
    (10, "Días de venta para pagar lo aplazado", 3, "0", "Regla de la cláusula: tercios, sin recargo"),
    (11, "Recargo del gota a gota por préstamo", bm.GOTA_A_GOTA_FLAT, PCT,
     "Enunciado del caso; 20,4% mensual en Cali (Martínez y Rivera-Acevedo 2019)"),
    (12, "Cuotas diarias del gota a gota (lunes a sábado)", 30, "0",
     "$4.000 diarios por cada $100.000 = 30 cuotas (Palomino-Martínez 2025, Palmira)"),
]

LEDGER_HEADER_ROW = 33
DAY0_ROW = LEDGER_HEADER_ROW + 1
FIRST_ROW = DAY0_ROW + 1
LAST_ROW = DAY0_ROW + LEDGER_DAYS

LEDGER_COLUMNS = [
    # (letter, header)
    ("A", "Día de venta"),
    ("B", "Lluvia (mm)"),
    ("C", "¿Rige la cláusula?"),
    ("D", "Ventas del día"),
    ("E", "Caja para pagar"),
    ("F", "Fiado del día"),
    ("G", "Cláusula: cuota de lo aplazado"),
    ("H", "Cláusula: paga al que fía"),
    ("I", "Cláusula: se aplaza hoy"),
    ("J", "Cláusula: faltante sin cubrir"),
    ("K", "Cláusula: saldo aplazado con quien fía"),
    ("L", "Cláusula: queda para la casa"),
    ("M", "Gota a gota: cuota de hoy"),
    ("N", "Gota a gota: préstamo nuevo"),
    ("O", "Gota a gota: saldo con el prestamista"),
    ("P", "Gota a gota: queda para la casa"),
]


def ledger_row_formulas(r):
    """Excel formulas for ledger row r (a selling day)."""
    p = r - 1
    return {
        "A": f"=A{p}+1",
        "C": f'=IF(B{r}>=$B$9,"Sí","No")',
        "D": f'=IF(C{r}="Sí",$B$21*$B$6,IF(B{r}>0,$B$21*$B$7,$B$21))',
        # On a rain day the vendor keeps up to $B$8 for the household; the rest goes to his debts.
        "E": f'=IF(C{r}="Sí",MAX(0,D{r}-$B$8),D{r})',
        "F": "=$B$4",
        # Each deferred amount is repaid in equal parts over the next $B$10 selling days.
        "G": f'=SUMIFS(I${DAY0_ROW}:I{p},A${DAY0_ROW}:A{p},">="&(A{r}-$B$10))/$B$10',
        "H": f"=MIN(F{r}+G{r},E{r})",
        "I": f'=IF(C{r}="Sí",F{r}+G{r}-H{r},0)',
        "J": f'=IF(C{r}="No",F{r}+G{r}-H{r},0)',
        "K": f"=K{p}+I{r}-G{r}",
        "L": f"=D{r}-H{r}",
        # Each loan is repaid with its flat charge in $B$12 equal daily installments.
        "M": f'=SUMIFS(N${DAY0_ROW}:N{p},A${DAY0_ROW}:A{p},">="&(A{r}-$B$12))*(1+$B$11)/$B$12',
        "N": f"=MAX(0,F{r}+M{r}-E{r})",
        "O": f"=O{p}+N{r}*(1+$B$11)-M{r}",
        "P": f"=D{r}-MIN(F{r}+M{r},E{r})",
    }


LEDGER_RANGE = f"{FIRST_ROW}:{LAST_ROW}"


def rng(col):
    return f"{col}{FIRST_ROW}:{col}{LAST_ROW}"


# Output rows: (row, label, clause formula, gota a gota formula, number format)
LEDGER_OUTPUTS = [
    (21, "Ventas de un día normal (fiado más ingreso)", "=B4+B5", None, PESOS),
    (24, "Días en que rigió la cláusula en el ejemplo", f'=COUNTIF({rng("C")},"Sí")',
     f'=COUNTIF({rng("C")},"Sí")', "0"),
    (25, "Recargo o interés que paga el vendedor", "=0", f"=SUM({rng('N')})*B11", PESOS),
    (26, "Ese costo como % del ingreso de un día normal", "=B25/$B$5", "=C25/$B$5", PCT),
    (27, "Préstamos de gota a gota que tuvo que tomar", "=0", f'=COUNTIF({rng("N")},">0")', "0"),
    (28, "De esos, préstamos que también pagan cuotas de otro gota a gota", "=0",
     f'=COUNTIFS({rng("N")},">0",{rng("M")},">0")', "0"),
    (29, "Queda para la casa en 30 días de venta (sin restar lo que sigue debiendo)", f"=SUM({rng('L')})",
     f"=SUM({rng('P')})", PESOS),
    (30, "Lo que sigue debiendo al día 30 (fuera del fiado del día)", f"=K{LAST_ROW}",
     f"=O{LAST_ROW}", PESOS),
    (31, "Casa menos deuda al día 30 (la diferencia es el interés)", "=B29-B30", "=C29-C30", PESOS),
    (32, "Mayor saldo aplazado que cargó quien fía", f"=MAX({rng('K')})", "=0", PESOS),
]


def style_header(cell):
    cell.font = WHITE_BOLD
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = BOX


def build_ledger_sheet(ws):
    ws.title = "Libro del vendedor"
    ws["A1"] = "Cláusula de lluvia del fiado: libro de un vendedor de carrito, día por día"
    ws["A1"].font = TITLE
    ws["A2"] = ("Celdas amarillas: supuestos que se pueden cambiar (también la lluvia de cada día). "
                "Todo lo verde y la tabla de abajo se calculan con fórmulas. [S] = supuesto propio sin dato.")
    ws["A3"] = "SUPUESTOS"
    ws["A3"].font = BOLD
    ws["C3"] = "Fuente o razón"
    ws["C3"].font = BOLD
    for row, label, value, fmt, source in LEDGER_INPUTS:
        ws[f"A{row}"] = label
        c = ws[f"B{row}"]
        c.value = value
        c.number_format = fmt
        c.fill = INPUT_FILL
        c.border = BOX
        ws[f"C{row}"] = source
    ws["A14"] = ("Lluvia de ejemplo (inventada para mostrar el cálculo): 3 días de disparo en 30, uno de ellos "
                 "una racha de 2 días, y 2 días de lluvia leve bajo el umbral. Dato real: 26 de 294 días de "
                 "venta con 5 mm o más, rachas de 1 o 2 días (IDEAM 0026055120, cálculo propio).")
    ws["A15"] = ("En los dos caminos, el día de lluvia el vendedor guarda para la casa hasta el tope de B8 y el "
                 "resto abona a su deuda. Camino con cláusula: lo que no alcanzó a pagar se paga en tercios, sin "
                 "recargo. Camino sin cláusula: ese hueco se cubre con un gota a gota al 20% en cuotas diarias.")

    ws["A20"] = "RESULTADOS"
    ws["A20"].font = BOLD
    ws["A23"] = "Indicador"
    ws["B23"] = "Con la cláusula"
    ws["C23"] = "Con gota a gota"
    for col in "ABC":
        style_header(ws[f"{col}23"])
    for row, label, f_clause, f_gota, fmt in LEDGER_OUTPUTS:
        ws[f"A{row}"] = label
        for col, formula in (("B", f_clause), ("C", f_gota)):
            if formula is None:
                continue
            c = ws[f"{col}{row}"]
            c.value = formula
            c.number_format = fmt
            c.fill = OUTPUT_FILL
            c.border = BOX

    for col, header in LEDGER_COLUMNS:
        style_header(ws[f"{col}{LEDGER_HEADER_ROW}"])
        ws[f"{col}{LEDGER_HEADER_ROW}"].value = header
    ws.row_dimensions[LEDGER_HEADER_ROW].height = 45
    # Day 0: starting state, nothing owed.
    for col, _ in LEDGER_COLUMNS:
        ws[f"{col}{DAY0_ROW}"] = 0 if col not in "C" else "Inicio"
    for r in range(FIRST_ROW, LAST_ROW + 1):
        day = r - DAY0_ROW
        rain = ws[f"B{r}"]
        rain.value = EXAMPLE_RAIN_MM.get(day, 0)
        rain.fill = INPUT_FILL
        rain.number_format = "0.0"
        for col, formula in ledger_row_formulas(r).items():
            c = ws[f"{col}{r}"]
            c.value = formula
            if col not in "AC":
                c.number_format = PESOS
        for col, _ in LEDGER_COLUMNS:
            ws[f"{col}{r}"].border = BOX
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    for col, _ in LEDGER_COLUMNS[3:]:
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = None


# ------------------------------------------------------------------------------------------
# Sheet 2: business numbers
# ------------------------------------------------------------------------------------------
BIZ_INPUTS = [
    (4, "Tarifa mensual por vendedor inscrito (la paga quien fía)", bm.BASE.fee, PESOS,
     "[S] sin precedente de precio; rango probado $2.000 a $6.000"),
    (5, "Parte de la tarifa que se cobra", bm.BASE.collection, PCT, "[S] 10% nunca se cobra"),
    (6, "Costo directo por vendedor al mes", bm.BASE.direct_cost, PESOS,
     "[S] mensajes de WhatsApp y tarjeta impresa"),
    (7, "Fijo: mi salario (1 SMMLV 2026)", bm.MINIMUM_WAGE_2026, PESOS, "Salario mínimo de 2026"),
    (8, "Fijo: transporte y campo", 200_000, PESOS, "[S]"),
    (9, "Fijo: celular y datos", 80_000, PESOS, "[S]"),
    (10, "Fijo: WhatsApp Business, servidor, herramientas", 150_000, PESOS, "[S]"),
    (11, "Fijo: contador", 350_000, PESOS, "[S]"),
    (12, "Inicial: registro de la SAS", 600_000, PESOS, "[S] sin cotización"),
    (13, "Inicial: concepto de un abogado", 3_000_000, PESOS, "[S] cotización pendiente"),
    (14, "Inicial: tarjetas y material del piloto", 300_000, PESOS, "[S]"),
    (15, "Inicial: capital para prestar", 0, PESOS, "La SAS no presta"),
    (16, "Vista de quien fía: fiado diario", bm.FIADO_UPPER_BOUND, PESOS, "Cota superior (ver hoja 1)"),
    (17, "Vista de quien fía: impago de lo aplazado", 0.03, PCT1,
     "[S] sin dato; referencia: mora del microcrédito formal 6,9% (El Tiempo con datos SFC, 2026)"),
    (18, "Vista de quien fía: margen bruto", 0.10, PCT, "[S] sin fuente"),
    (19, "Días de disparo al año", bm.TRIGGER_DAYS_PER_YEAR, "0", "IDEAM 0026055120, cálculo propio"),
    (20, "Días promedio que tarda en volver lo aplazado", 2.5, "0.0", "[S] tercios más domingos"),
    (21, "Costo del dinero de quien fía (efectivo anual)", 0.20, PCT, "[S]"),
    (22, "Días de venta al mes", bm.SELLING_DAYS_PER_MONTH, "0", "Lunes a sábado"),
]

MONTHS = 36
TABLE_HEADER_ROW = 45
TABLE_FIRST = TABLE_HEADER_ROW + 1
TABLE_LAST = TABLE_HEADER_ROW + MONTHS
TABLE_COLUMNS = ["Mes", "Minoristas", "Vendedores por minorista", "Vendedores inscritos",
                 "Vendedores que pagan", "Facturado", "Cobrado", "Contribución", "Costo fijo",
                 "Resultado del mes", "Acumulado (con inversión)", "¿Cubre el costo fijo?"]


def tcol(letter):
    return f"{letter}{TABLE_FIRST}:{letter}{TABLE_LAST}"


BIZ_OUTPUTS = [
    (26, "Inversión inicial", "=SUM(B12:B15)", PESOS),
    (27, "Costo fijo mensual", "=SUM(B7:B11)", PESOS),
    (28, "Contribución por vendedor al mes", "=B4*B5-B6", PESOS),
    (29, "Vendedores que pagan para el equilibrio", "=B27/B28", "#,##0"),
    (30, "Mes del equilibrio", f'=IFERROR(INDEX({tcol("A")},MATCH(1,{tcol("L")},0)),"no llega en {MONTHS} meses")',
     "0"),
    (31, "Caja necesaria (inversión más pérdidas hasta el punto más bajo)",
     f"=-MIN(MIN({tcol('K')}),-B26)", PESOS),
    (32, "Equipo de $10 millones al mes: vendedores para el equilibrio", "=10000000/B28", "#,##0"),
    (34, "Quien fía: aplazado al año por vendedor", "=B19*B16", PESOS),
    (35, "Quien fía: costo del plazo al año", "=B34*((1+B21)^(B20/365)-1)", PESOS),
    (36, "Quien fía: impago de lo aplazado al año", "=B34*B17", PESOS),
    (37, "Quien fía: tarifa al año", "=B4*12", PESOS),
    (38, "Quien fía: margen que le deja un vendedor al año", "=B16*B22*12*B18", PESOS),
    (39, "Se paga sola si le evita perder esta parte de sus vendedores al año", "=(B35+B36+B37)/B38", PCT1),
]
MILESTONE_ROWS = {3: 41, 6: 42, 12: 43}


def build_business_sheet(ws):
    ws.title = "Números del negocio"
    ws["A1"] = "Cláusula de lluvia del fiado: números de la SAS (caso base)"
    ws["A1"].font = TITLE
    ws["A2"] = ("Celdas amarillas: supuestos, incluidas las columnas de minoristas y vendedores por minorista "
                "de la tabla mensual. Lo verde se calcula. Mismos supuestos que models/business_model.py.")
    ws["A3"] = "SUPUESTOS"
    ws["A3"].font = BOLD
    ws["C3"] = "Fuente o razón"
    ws["C3"].font = BOLD
    for row, label, value, fmt, source in BIZ_INPUTS:
        ws[f"A{row}"] = label
        c = ws[f"B{row}"]
        c.value = value
        c.number_format = fmt
        c.fill = INPUT_FILL
        c.border = BOX
        ws[f"C{row}"] = source
    ws["A25"] = "RESULTADOS"
    ws["A25"].font = BOLD
    for row, label, formula, fmt in BIZ_OUTPUTS:
        ws[f"A{row}"] = label
        c = ws[f"B{row}"]
        c.value = formula
        c.number_format = fmt
        c.fill = OUTPUT_FILL
        c.border = BOX
    ws["A40"] = "Usuarios e ingresos"
    for col, head in zip("BCDE", ["Inscritos", "Pagan", "Facturado al mes", "Contribución al mes"]):
        style_header(ws[f"{col}40"])
        ws[f"{col}40"].value = head
    style_header(ws["A40"])
    for month, row in MILESTONE_ROWS.items():
        tr = TABLE_HEADER_ROW + month
        ws[f"A{row}"] = f"Mes {month}"
        for col, src, fmt in (("B", "D", "#,##0"), ("C", "E", "#,##0"), ("D", "F", PESOS), ("E", "H", PESOS)):
            c = ws[f"{col}{row}"]
            c.value = f"={src}{tr}"
            c.number_format = fmt
            c.fill = OUTPUT_FILL
            c.border = BOX

    for i, head in enumerate(TABLE_COLUMNS):
        col = get_column_letter(i + 1)
        ws[f"{col}{TABLE_HEADER_ROW}"] = head
        style_header(ws[f"{col}{TABLE_HEADER_ROW}"])
    ws.row_dimensions[TABLE_HEADER_ROW].height = 45
    for m in range(1, MONTHS + 1):
        r = TABLE_HEADER_ROW + m
        p = r - 1
        ws[f"A{r}"] = m
        ws[f"B{r}"] = bm.base_retailers(m)
        ws[f"C{r}"] = bm.base_vendors_per_retailer(m)
        for col in "BC":
            ws[f"{col}{r}"].fill = INPUT_FILL
        ws[f"D{r}"] = f"=B{r}*C{r}"
        # First month free for each retailer: only vendors already enrolled last month pay.
        ws[f"E{r}"] = "=0" if m == 1 else f"=MIN(D{p},D{r})"
        ws[f"F{r}"] = f"=E{r}*$B$4"
        ws[f"G{r}"] = f"=F{r}*$B$5"
        ws[f"H{r}"] = f"=E{r}*$B$28"
        ws[f"I{r}"] = "=$B$27"
        ws[f"J{r}"] = f"=H{r}-I{r}"
        ws[f"K{r}"] = f"=-$B$26+J{r}" if m == 1 else f"=K{p}+J{r}"
        ws[f"L{r}"] = f"=IF(J{r}>=0,1,0)"
        for col in "FGHIJK":
            ws[f"{col}{r}"].number_format = PESOS
        for i in range(len(TABLE_COLUMNS)):
            ws[f"{get_column_letter(i + 1)}{r}"].border = BOX
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    for col in "DEFGHIJKL":
        ws.column_dimensions[col].width = 15


# ------------------------------------------------------------------------------------------
# Independent Python mirror of the ledger, to check the spreadsheet
# ------------------------------------------------------------------------------------------
def mirror_ledger():
    inp = {row: value for row, _, value, _, _ in LEDGER_INPUTS}
    fiado, income, rain_share, light_share, home_floor = inp[4], inp[5], inp[6], inp[7], inp[8]
    threshold, clause_days, gota_flat, gota_n = inp[9], inp[10], inp[11], inp[12]
    normal = fiado + income
    deferred, loans = {0: 0.0}, {0: 0.0}
    k_bal = o_bal = 0.0
    out = {"home_c": 0.0, "home_g": 0.0, "interest": 0.0, "k_max": 0.0}
    for t in range(1, LEDGER_DAYS + 1):
        mm = EXAMPLE_RAIN_MM.get(t, 0)
        trig = mm >= threshold
        sales = normal * rain_share if trig else (normal * light_share if mm > 0 else normal)
        cash = max(0.0, sales - home_floor) if trig else sales
        g = sum(v for d, v in deferred.items() if d >= t - clause_days) / clause_days
        h = min(fiado + g, cash)
        deferred[t] = fiado + g - h if trig else 0.0
        k_bal += deferred[t] - g
        out["home_c"] += sales - h
        out["k_max"] = max(out["k_max"], k_bal)
        m = sum(v for d, v in loans.items() if d >= t - gota_n) * (1 + gota_flat) / gota_n
        loans[t] = max(0.0, fiado + m - cash)
        o_bal += loans[t] * (1 + gota_flat) - m
        out["home_g"] += sales - min(fiado + m, cash)
    out["interest"] = sum(loans.values()) * gota_flat
    out["k_end"], out["o_end"] = k_bal, o_bal
    return out


def evaluate(path):
    """Evaluate every formula with the `formulas` package; returns {(sheet, cell): value}."""
    import formulas  # noqa: PLC0415 (optional dependency, only needed for the check)

    model = formulas.ExcelModel().loads(path).finish()
    solution = model.calculate()
    values = {}
    for key, ranges in solution.items():
        # keys look like "'[prototipo.xlsx]LIBRO DEL VENDEDOR'!B24"
        if "!" not in key or ":" in key.split("!")[1]:
            continue
        sheet, cell = key.rsplit("!", 1)
        sheet = sheet.strip("'").split("]", 1)[1].upper()
        value = ranges.value[0][0]
        values[(sheet, cell.upper())] = value
    return values


def fmt_pesos(v):
    s = f"{round(v):,}".replace(",", ".")
    return f"-${s[1:]}" if s.startswith("-") else f"${s}"


def fmt_value(v, fmt):
    if isinstance(v, str):
        return v
    if fmt == PESOS:
        return fmt_pesos(v)
    if fmt == PCT:
        return f"{v * 100:.0f}%"
    if fmt == PCT1:
        return f"{v * 100:.1f}%".replace(".", ",")
    if fmt == "0.0":
        return f"{v:.1f}".replace(".", ",")
    return f"{round(v):,}".replace(",", ".")


def write_snapshot(values):
    s1, s2 = "LIBRO DEL VENDEDOR", "NÚMEROS DEL NEGOCIO"

    def v1(cell):
        return values[(s1, cell)]

    def v2(cell):
        return values[(s2, cell)]

    parts = ["""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><title>Prototipo: libro del vendedor</title>
<style>
.proto-snap { font-family: "Helvetica Neue", Helvetica, Arial, sans-serif; font-size: 7.4pt; color: #111; }
.proto-snap h3 { font-size: 9pt; margin: 8pt 0 3pt 0; }
.proto-snap p { margin: 2pt 0 4pt 0; }
.proto-snap table { border-collapse: collapse; width: 100%; margin-bottom: 6pt; }
.proto-snap th, .proto-snap td { border: 0.5pt solid #b8b8b8; padding: 1.5pt 3pt; text-align: right; }
.proto-snap th { background: #1f4e79; color: #fff; font-weight: 600; text-align: center; }
.proto-snap td.l { text-align: left; }
.proto-snap td.in { background: #fff2cc; }
.proto-snap td.out { background: #e2efda; }
.proto-snap tr.rain td { background: #ddebf7; }
</style></head><body><section class="proto-snap">
"""]
    parts.append("<h3>Hoja 1. Libro de un vendedor de carrito: con la cláusula y con gota a gota</h3>")
    parts.append("<p>Supuestos (amarillo) y resultados (verde) de 30 días de venta. La lluvia del ejemplo es "
                 "inventada para mostrar el cálculo; el dato real es 26 de 294 días de venta con 5 mm o más "
                 "(IDEAM 0026055120, cálculo propio).</p>")
    parts.append("<table><tr><th>Supuesto</th><th>Valor</th></tr>")
    for row, label, _, fmt, _ in LEDGER_INPUTS:
        parts.append(f'<tr><td class="l">{html.escape(label)}</td>'
                     f'<td class="in">{fmt_value(v1(f"B{row}"), fmt)}</td></tr>')
    parts.append("</table>")
    parts.append("<table><tr><th>Resultado en 30 días de venta</th><th>Con la cláusula</th>"
                 "<th>Con gota a gota</th></tr>")
    for row, label, _, f_gota, fmt in LEDGER_OUTPUTS[1:]:
        gota = fmt_value(v1(f"C{row}"), fmt) if f_gota else ""
        parts.append(f'<tr><td class="l">{html.escape(label)}</td>'
                     f'<td class="out">{fmt_value(v1(f"B{row}"), fmt)}</td><td class="out">{gota}</td></tr>')
    parts.append("</table>")

    shown = ["A", "B", "C", "D", "G", "H", "I", "K", "L", "M", "N", "O", "P"]
    heads = dict(LEDGER_COLUMNS)
    parts.append("<table><tr>" + "".join(f"<th>{html.escape(heads[c])}</th>" for c in shown) + "</tr>")
    for r in range(FIRST_ROW, LAST_ROW + 1):
        day = r - DAY0_ROW
        if day not in SNAPSHOT_DAYS:
            continue
        cls = ' class="rain"' if v1(f"C{r}") == "Sí" else ""
        cells = []
        for c in shown:
            val = v1(f"{c}{r}")
            if c == "A":
                cells.append(f"<td>{int(val)}</td>")
            elif c == "B":
                cells.append(f"<td>{fmt_value(val, '0.0')}</td>")
            elif c == "C":
                cells.append(f"<td>{val}</td>")
            else:
                cells.append(f"<td>{fmt_pesos(val)}</td>")
        parts.append(f"<tr{cls}>" + "".join(cells) + "</tr>")
    parts.append("</table>")
    parts.append("<p>Fiado de $62.000 todos los días. Se muestran 15 de los 30 días (la hoja los tiene todos). "
                 "En azul, los días en que rige la cláusula; los días 9 y 21 llueve poco, bajan las ventas "
                 "y la cláusula no se activa.</p>")

    parts.append("<h3>Hoja 2. Números de la SAS (caso base)</h3>")
    parts.append("<table><tr><th>Resultado</th><th>Valor</th></tr>")
    for row, label, _, fmt in BIZ_OUTPUTS:
        parts.append(f'<tr><td class="l">{html.escape(label)}</td>'
                     f'<td class="out">{fmt_value(v2(f"B{row}"), fmt)}</td></tr>')
    parts.append("</table>")
    parts.append("<table><tr><th>Mes</th><th>Inscritos</th><th>Pagan</th><th>Facturado al mes</th>"
                 "<th>Contribución al mes</th></tr>")
    for month, row in MILESTONE_ROWS.items():
        parts.append(f"<tr><td>{month}</td><td>{fmt_value(v2(f'B{row}'), '#,##0')}</td>"
                     f"<td>{fmt_value(v2(f'C{row}'), '#,##0')}</td><td>{fmt_pesos(v2(f'D{row}'))}</td>"
                     f"<td>{fmt_pesos(v2(f'E{row}'))}</td></tr>")
    parts.append("</table>")
    parts.append("<p>Generado desde report/prototipo.xlsx con models/build_prototype.py (fórmulas evaluadas con "
                 "la librería formulas y verificadas contra models/business_model.py). Los supuestos marcados [S] "
                 "y sus fuentes están en la hoja.</p>")
    parts.append("</section></body></html>\n")
    with open(HTML_PATH, "w", encoding="utf-8") as fh:
        fh.write("\n".join(parts))


def check(values):
    s1, s2 = "LIBRO DEL VENDEDOR", "NÚMEROS DEL NEGOCIO"
    mirror = mirror_ledger()
    pairs = [
        ("home with clause", values[(s1, "B29")], mirror["home_c"]),
        ("home with gota a gota", values[(s1, "C29")], mirror["home_g"]),
        ("gota a gota interest", values[(s1, "C25")], mirror["interest"]),
        ("deferred balance day 30", values[(s1, "B30")], mirror["k_end"]),
        ("gota a gota balance day 30", values[(s1, "C30")], mirror["o_end"]),
        ("max deferred balance", values[(s1, "B32")], mirror["k_max"]),
    ]
    sim = bm.simulate(bm.BASE, horizon=MONTHS)
    pairs += [
        ("initial investment", values[(s2, "B26")], sum(bm.INITIAL_INVESTMENT.values())),
        ("fixed monthly", values[(s2, "B27")], sim["fixed"]),
        ("contribution per vendor", values[(s2, "B28")], sim["unit"]),
        ("break-even vendors", values[(s2, "B29")], sim["break_even_vendors"]),
        ("break-even month", values[(s2, "B30")], sim["break_even_month"]),
        ("cash needed", values[(s2, "B31")], sim["cash_needed"]),
        ("month 12 contribution", values[(s2, "E43")], sim["rows"][11]["contribution"]),
    ]
    payer = bm.payer_view(bm.FIADO_UPPER_BOUND, 0.03)
    pairs.append(("payer churn to pay off", values[(s2, "B39")], payer["churn_to_pay_off"]))
    ok = True
    for name, got, want in pairs:
        good = abs(float(got) - float(want)) < 0.5
        ok &= good
        print(f"  {'ok ' if good else 'BAD'} {name:28s} sheet {float(got):>14,.2f}  python {float(want):>14,.2f}")
    return ok


def main():
    wb = Workbook()
    build_ledger_sheet(wb.active)
    build_business_sheet(wb.create_sheet())
    wb.save(XLSX_PATH)
    print(f"wrote {XLSX_PATH}")
    values = evaluate(XLSX_PATH)
    print("check against the Python mirror and business_model.py:")
    if not check(values):
        sys.exit("spreadsheet and Python disagree")
    write_snapshot(values)
    print(f"wrote {HTML_PATH}")
    s1 = "LIBRO DEL VENDEDOR"
    for row, label, *_ in LEDGER_OUTPUTS[1:]:
        print(f"  {label:66s} clause {values[(s1, f'B{row}')]!s:>12}  gota {values[(s1, f'C{row}')]!s:>12}")


if __name__ == "__main__":
    main()
