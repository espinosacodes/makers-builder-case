"""Build Defensa_Builder_Case.pptx from pitch/slides.md (titles, sources, notes parsed from the md)."""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

HERE = Path(__file__).resolve().parent
MD = HERE.parent / "slides.md"
OUT = HERE / "Defensa_Builder_Case.pptx"

INK = RGBColor(0x11, 0x11, 0x11)
ACCENT = RGBColor(0x15, 0x58, 0xB0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
W, H = 13.333, 7.5
M = 0.625  # 60 px margin at 96 dpi
CW = W - 2 * M


# ---------- parse slides.md ----------
def parse_md():
    text = MD.read_text(encoding="utf-8")
    parts = re.split(r"^## Diapositiva (\d+)[^\n]*\n", text, flags=re.M)
    slides = {}
    for i in range(1, len(parts), 2):
        n, body = int(parts[i]), parts[i + 1]
        title = re.search(r"^\*\*Título \(el mensaje\):\*\* (.+)$", body, re.M).group(1).strip()
        src = re.search(r"^\*\*Línea de fuente \(12 pt\):\*\* (.+)$", body, re.M).group(1).strip()
        if src.startswith(("Ninguna", "Sin fuente")):
            src = ""
        top = re.search(r"^\*\*Línea superior \(12 pt\):\*\* (.+)$", body, re.M)
        notes_block = body.split("**Notas del orador", 1)[1].split("\n", 1)[1]
        paras, cur = [], []
        for line in notes_block.splitlines():
            if line.startswith(">"):
                content = line[1:].strip()
                if content:
                    cur.append(content)
                elif cur:
                    paras.append(" ".join(cur)); cur = []
            elif cur:
                paras.append(" ".join(cur)); cur = []
        if cur:
            paras.append(" ".join(cur))
        stage = re.search(r"^\*(\[.+\])\*$", notes_block, re.M)
        notes = "\n\n".join(paras) + ("\n\n" + stage.group(1) if stage else "")
        slides[n] = dict(title=title, src=src, top=top.group(1).strip() if top else "", notes=notes)
    return slides


# ---------- helpers ----------
def style_run(run, size, bold=False, color=INK, strike=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)
    if strike:
        rpr.set("strike", "sngStrike")


def textbox(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """paras: list of paragraphs; each paragraph is a list of (text, size, bold, color[, strike])."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    fill_frame(tf, paras, align)
    return tb


def fill_frame(tf, paras, align):
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for spec in runs:
            text, size, bold, color = spec[:4]
            strike = spec[4] if len(spec) > 4 else False
            style_run(p.add_run(), size, bold, color, strike)
            p.runs[-1].text = text


def rect(slide, x, y, w, h, fill=None, line=INK, line_pt=2, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_pt)
    s.shadow.inherit = False
    return s


def line(slide, x1, y1, x2, y2, pt=1, color=INK):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(pt)
    return c


def base(prs, n, data, title_size=32, title=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid(); bg.fore_color.rgb = WHITE
    if title:
        textbox(s, M, 0.55 if not data["top"] else 0.75, CW, 1.3,
                [[(data["title"], title_size, True, INK)]])
    if data["src"]:
        textbox(s, M, 6.55, CW - 1.2, 0.6, [[(data["src"], 12, False, INK)]], anchor=MSO_ANCHOR.BOTTOM)
    textbox(s, W - M - 1.0, 6.85, 1.0, 0.3, [[(str(n), 12, False, INK)]], align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.BOTTOM)
    tf = s.notes_slide.notes_text_frame
    tf.text = ""
    fill_frame(tf, [[(p, 12, False, INK)] for p in data["notes"].split("\n\n")], PP_ALIGN.LEFT)
    return s


# ---------- slides ----------
def s1(prs, d):
    s = base(prs, 1, d)
    textbox(s, M, 0.35, CW, 0.3, [[(d["top"], 12, False, INK)]])
    bw, bh, by = 3.4, 2.2, 2.55
    gap = (CW - 3 * bw) / 2
    xs = [M + i * (bw + gap) for i in range(3)]
    for x in xs:
        rect(s, x, by, bw, bh)
    textbox(s, xs[0], by, bw, bh, [[("Lluvia ≥", 28, False, INK)], [("5 mm", 72, True, ACCENT)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, xs[1] + 0.15, by, bw - 0.3, bh, [[("Guarda hasta $15.000", 28, False, INK)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, xs[2] + 0.15, by, bw - 0.3, bh, [[("Tercios, 3 días, sin recargo", 28, False, INK)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(2):
        ax = xs[i] + bw + (gap - 0.6) / 2
        rect(s, ax, by + bh / 2 - 0.2, 0.6, 0.4, fill=INK, line=None, shape=MSO_SHAPE.RIGHT_ARROW)
    textbox(s, M, 5.3, CW, 0.7, [[("No presto. No aseguro.", 28, True, INK)]], align=PP_ALIGN.CENTER)


def s2(prs, d):
    s = base(prs, 2, d, 36)
    full = CW  # $4.000 maps to full content width; axis starts at zero
    textbox(s, M, 1.85, CW, 0.5, [[("Legal: $3.436", 26, False, INK)]])
    rect(s, M, 2.4, full * 3436 / 4000, 0.55, fill=INK, line=None)
    textbox(s, M, 3.1, CW, 0.5, [[("Gota a gota: $4.000", 26, False, INK)]])
    rect(s, M, 3.65, full, 0.55, fill=ACCENT, line=None)
    textbox(s, M, 4.3, CW, 0.5, [[("diarios por $100.000", 24, False, INK)]])
    textbox(s, M, 5.2, CW, 0.8, [[("87,2%", 40, True, INK), (" del formal: cuota mensual", 28, False, INK)]],
            anchor=MSO_ANCHOR.BOTTOM)


def s3(prs, d):
    s = base(prs, 3, d, 36)
    textbox(s, M, 2.0, CW, 0.5, [[("Colaborador:", 24, False, INK)]])
    textbox(s, M, 2.7, CW, 2.0, [
        [("“…", 36, False, INK), ("llovió", 36, True, ACCENT), (", entonces nadie va…”", 36, False, INK)],
        [("“Y ahí es cuando entran con un gota a gota.”", 36, False, INK)],
    ])


def s4(prs, d):
    s = base(prs, 4, d, 32)
    lw = 5.0
    rw = CW - lw
    textbox(s, M, 2.3, lw, 1.6, [[("61,8%", 96, True, ACCENT)]], anchor=MSO_ANCHOR.BOTTOM)
    textbox(s, M, 4.0, lw - 0.3, 1.2, [[("ambulantes con crédito: gota a gota", 26, False, INK)]])
    textbox(s, M + lw, 2.3, rw, 1.6, [[("26 de 294", 96, True, ACCENT)]], anchor=MSO_ANCHOR.BOTTOM)
    textbox(s, M + lw, 4.0, rw, 1.2, [[("días de venta con lluvia", 26, False, INK)]])


def s5(prs, d):
    s = base(prs, 5, d, 36)
    rows = [
        ("Préstamo diario", [("$5.857 a -$26.143", 48, True, ACCENT)], [("al año", 24, False, INK)]),
        ("Seguro de lluvia", [("43%", 48, True, ACCENT)], [("del ingreso, cada día", 24, False, INK)]),
    ]
    lw, rh = 4.6, 1.75
    gf = s.shapes.add_table(2, 2, Inches(M), Inches(2.1), Inches(CW), Inches(2 * rh))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tbl.columns[0].width = Inches(lw)
    tbl.columns[1].width = Inches(CW - lw)
    for r, (name, num, sub) in enumerate(rows):
        tbl.rows[r].height = Inches(rh)
        for c in range(2):
            cell = tbl.cell(r, c)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.1)
            set_borders(cell, top=(r == 0), bottom=True)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
        fill_frame(tbl.cell(r, 0).text_frame, [[(name, 32, False, INK, True)]], PP_ALIGN.LEFT)
        fill_frame(tbl.cell(r, 1).text_frame, [num, sub], PP_ALIGN.LEFT)


def set_borders(cell, top=False, bottom=True):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag, on in (("a:lnL", False), ("a:lnR", False), ("a:lnT", top), ("a:lnB", bottom)):
        ln = tcPr.makeelement(qn(tag), {"w": str(Pt(1)) if on else "0"})
        if on:
            sf = ln.makeelement(qn("a:solidFill"), {})
            clr = sf.makeelement(qn("a:srgbClr"), {"val": "111111"})
            sf.append(clr); ln.append(sf)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.append(ln)


def s6(prs, d):
    s = base(prs, 6, d, 36)
    top, bot = 2.0, 6.4
    cx, cy = W / 2, (top + bot) / 2
    line(s, cx, top, cx, bot, pt=1)
    line(s, M, cy, W - M, cy, pt=1)
    r = 0.25
    rect(s, cx - r, cy - r, 2 * r, 2 * r, fill=ACCENT, line=None, shape=MSO_SHAPE.OVAL)
    qw, qh = CW / 2 - 0.4, (bot - top) / 2 - 0.2
    corners = ["No presto", "Solo el día sin venta", "Su minorista de siempre", "Plazo por dato"]
    pos = [(M, top), (cx + 0.4, top), (M, cy + 0.2), (cx + 0.4, cy + 0.2)]
    for text, (x, y) in zip(corners, pos):
        textbox(s, x, y, qw, qh, [[(text, 28, False, INK)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def s7(prs, d):
    s = base(prs, 7, d, 32)
    cols = [("Vendedor", "0", "pendiente"), ("Quien fía", "2,5%", "de su margen"), ("Mi empresa", "Mes 10", "sin salario")]
    cw = CW / 3
    for i, (head, big, sub) in enumerate(cols):
        x = M + i * cw
        textbox(s, x, 2.2, cw, 0.5, [[(head, 24, True, INK)]], align=PP_ALIGN.CENTER)
        textbox(s, x, 2.8, cw, 1.4, [[(big, 80, True, ACCENT)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x, 4.3, cw, 0.5, [[(sub, 24, False, INK)]], align=PP_ALIGN.CENTER)


def s8(prs, d):
    s = base(prs, 8, d, 32)
    pts = [("Antes:", "6 vendedores"), ("Semana 1:", "gratis"), ("Lluvia:", "WhatsApp"), ("3 días:", "tercios"),
           ("Mes:", "Nequi")]
    x0, x1, y = M + 1.2, W - M - 1.2, 3.3
    line(s, x0, y, x1, y, pt=2)
    step = (x1 - x0) / 4
    r = 0.2
    for i, (a, b) in enumerate(pts):
        x = x0 + i * step
        rect(s, x - r, y - r, 2 * r, 2 * r, fill=ACCENT if i in (0, 2) else INK, line=None, shape=MSO_SHAPE.OVAL)
        textbox(s, x - 1.3, y + 0.45, 2.6, 1.2, [[(a, 24, True, INK)], [(b, 24, False, INK)]],
                align=PP_ALIGN.CENTER)


def s9(prs, d):
    s = base(prs, 9, d, 32)
    textbox(s, M, 2.0, CW, 2.3, [[("3 de 6", 140, True, ACCENT)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, M, 4.6, CW, 0.6, [[("Caja del minorista un día de lluvia: $562.000 con 12 vendedores", 24, False, INK)]],
            align=PP_ALIGN.CENTER)


def s10(prs, d):
    s = base(prs, 10, d, title=False)
    phrase = "Le falta que un día de lluvia no cuente como quedar mal."
    a, b = phrase.split("un día de lluvia")
    textbox(s, M, 1.6, CW, 4.2, [
        [(d["title"], 44, True, INK)],
        [(a, 36, False, INK), ("un día de lluvia", 36, True, ACCENT), (b, 36, False, INK)],
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def main():
    data = parse_md()
    assert sorted(data) == list(range(1, 11)), sorted(data)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for n, fn in enumerate([s1, s2, s3, s4, s5, s6, s7, s8, s9, s10], start=1):
        fn(prs, data[n])
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
