"""Build Builder_Case_Docs.docx and informe-docs.txt from informe-docs.md.

The .docx uses Arial 11 black text, bold black headings, 1.15 line spacing,
US Letter with 1 inch margins, thin-bordered 9pt tables, no hyperlinks, no
colors, and a page break before ANEXOS. The .txt is a plain fallback for paste.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
MD = HERE / "informe-docs.md"
DOCX_OUT = HERE / "Builder_Case_Docs.docx"
TXT_OUT = HERE / "informe-docs.txt"
ANNEX_MARKER = "ANEXOS"
FONT = "Arial"
BLACK = RGBColor(0, 0, 0)

# Column widths in inches (6.5 in of usable width), keyed by the first header cell.
TABLE_WIDTHS = {
    "Hallazgo": [2.0, 2.2, 2.3],
    "Cuándo y quién": [1.35, 3.55, 1.6],
    "Pregunta y respuesta": [2.9, 3.6],
}


def set_run_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text, size=11, force_bold=False):
    """Add text with **bold** and *italic* spans as separate runs."""
    tokens = re.split(r"(\*\*.+?\*\*|(?<![*\w])\*(?!\s).+?(?<!\s)\*(?![*\w]))", text)
    for tok in tokens:
        if not tok:
            continue
        bold, italic = force_bold, False
        if tok.startswith("**") and tok.endswith("**"):
            tok, bold = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 1:
            tok, italic = tok[1:-1], True
        set_run_font(paragraph.add_run(tok), size=size, bold=bold, italic=italic)


def format_paragraph(paragraph, before=0, after=6, line=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_cell_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tbl_pr.append(borders)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def style_document(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    for name in ("Normal", "List Bullet", "Heading 1", "Heading 2", "Heading 3", "Title"):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.color.rgb = BLACK
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), FONT)
        # Drop theme fonts so Word does not fall back to Calibri or Cambria.
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
    doc.styles["Normal"].font.size = Pt(11)


def build_docx(md):
    doc = Document()
    style_document(doc)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip() == ANNEX_MARKER:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            add_inline(p, ANNEX_MARKER, size=14, force_bold=True)
            format_paragraph(p)
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            size = {1: 16, 2: 14, 3: 12}[level]
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, m.group(2), size=size, force_bold=True)
            format_paragraph(p, before=0 if level == 1 else 12, after=6)
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if line.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = split_row(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.autofit = False
            set_cell_borders(table)
            widths = TABLE_WIDTHS.get(header[0], [6.5 / len(header)] * len(header))
            for r_idx, cells in enumerate([header] + rows):
                for c_idx, text in enumerate(cells):
                    cell = table.cell(r_idx, c_idx)
                    cell.width = Inches(widths[c_idx])
                    p = cell.paragraphs[0]
                    add_inline(p, text, size=9, force_bold=(r_idx == 0))
                    format_paragraph(p, after=0)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, re.sub(r"^\s*[-*]\s+", "", lines[i]))
                format_paragraph(p, after=4)
                i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line.strip())
        format_paragraph(p)
        i += 1
    doc.save(DOCX_OUT)


def build_txt(md):
    out = []
    for line in md.splitlines():
        if re.match(r"^\s*\|[\s:|-]+\|\s*$", line):
            continue
        if line.lstrip().startswith("|"):
            line = " | ".join(split_row(line))
        line = re.sub(r"^#{1,3}\s+", "", line)
        line = re.sub(r"^\s*[-*]\s+", "• ", line)
        line = line.replace("**", "")
        line = re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])", r"\1", line)
        out.append(line)
    TXT_OUT.write_text("\n".join(out).strip() + "\n", encoding="utf-8")


def main():
    md = MD.read_text(encoding="utf-8")
    build_docx(md)
    build_txt(md)
    print(f"wrote {DOCX_OUT.name} and {TXT_OUT.name}")


if __name__ == "__main__":
    main()
